# Copyright (c) 2026, Cédric Nguendap Bedjama and contributors
# For license information, please see license.txt

import re

import frappe
from frappe import _
from frappe.utils import flt, today

from udshed.grade_calculation import (
    combinaison_detectee,
    combinaison_type_ue,
    get_formula,
    get_seuil_validation,
)

TYPE_CC = "Controlle Continue (CC)"
TYPE_NORMALE = "Examen de session normal"
TYPE_RATTRAPAGE = "Examen de rattrapage"

TYPES_EVAL = [TYPE_CC, TYPE_NORMALE, TYPE_RATTRAPAGE]

NOTE_MAX = 20


def _cycle_pour_niveau(niveau_label):
    """Cycle (Licence / BTS / Master) déduit du libellé d'un niveau.

    Les libellés de niveau sont normalisés par la grille de la filière
    (ex : « Licence 1 », « Master 2 », « BTS 1 »). Repli sur Licence.
    """
    label = (niveau_label or "").lower()
    if label.startswith("master"):
        return "Master"
    if label.startswith("bts"):
        return "BTS"
    return "Licence"


def _serialiser_formule(formula, cycle, combinaison):
    """Représentation lisible d'une Grade Formula pour la page de saisie.

    Source de vérité des poids : les composantes de la formule active.
    Aucun poids n'est saisi manuellement dans la saisie des notes.
    """
    composantes = [
        {"composante": c.composante, "pourcentage": float(c.pourcentage or 0)}
        for c in formula.components
        if (c.pourcentage or 0) > 0
    ]
    return {
        "nom": formula.name,
        "cycle": cycle,
        "combinaison": combinaison,
        "composantes": composantes,
        "methode_calcul_cc": formula.methode_calcul_cc,
        "nb_meilleures_notes_cc": formula.nb_meilleures_notes_cc,
        "seuil_validation": formula.seuil_validation,
        "methode_arrondi": formula.methode_arrondi,
    }


def _formule_pour_ue(teaching_unit, filiere=None, niveau=None):
    """Formule active attendue pour une UE (cycle du niveau + combinaison du type d'UE).

    Returns:
        dict | None: formule sérialisée, ou None si aucune formule active.
    """
    type_ue = frappe.db.get_value("Teaching Unit", teaching_unit, "type_ue") or "Sans TP"
    cycle = _cycle_pour_niveau(niveau or "")
    combinaison = combinaison_type_ue(type_ue)
    formula = get_formula(cycle, combinaison)
    if not formula:
        return None
    return _serialiser_formule(formula, cycle, combinaison)


def _formule_detectee(teaching_unit, filiere, niveau, notes, students):
    """Formule correspondant à la combinaison détectée des notes réellement saisies.

    La combinaison est l'union des évaluations renseignées dans la session
    normale (CC, CCTP, EXAMTP, EXAM). Sans aucune note saisie, on retombe sur
    la formule attendue du type d'UE.

    Returns:
        dict | None: formule sérialisée, ou None si aucune formule active.
    """
    remplis = set()
    for s in students or []:
        ex = (notes or {}).get("Examen", {}).get(s["student"], {})
        # Drapeaux *saisi : le stockage Currency ramène les champs non
        # saisis à 0.0, ils ne comptent donc jamais comme remplis.
        if ex.get("cc_saisi"):
            remplis.add("Controle Continu(CC)")
        if ex.get("cctp_saisi"):
            remplis.add("Controle Continu Travaux Pratiques(CCTP)")
        if ex.get("examtp_saisi"):
            remplis.add("Examen Travaux Pratiques(EXAMTP)")
        if ex.get("examen_saisi") or ex.get("rattrapage_saisi"):
            remplis.add("Examen")
    if remplis:
        cycle = _cycle_pour_niveau(niveau or "")
        combinaison = combinaison_detectee(remplis)
        formula = get_formula(cycle, combinaison)
        if formula:
            return _serialiser_formule(formula, cycle, combinaison)
    return _formule_pour_ue(teaching_unit, filiere, niveau)


def _get_enseignant_courant():
    """Enseignant (Teacher) lié à l'utilisateur connecté, ou None."""
    if frappe.session.user == "Administrator":
        return None
    return frappe.db.get_value("Teacher", {"email": frappe.session.user}, "name")


def _get_enseignant_full_name(enseignant):
    if not enseignant:
        return ""
    return frappe.db.get_value("Teacher", enseignant, "full_name") or enseignant


def _semestre_effectif(args):
    """Semestre demandé, ou celui de l'UE si le client n'en a pas transmis."""
    semestre = args.get("semestre")
    if semestre:
        return semestre
    return frappe.db.get_value("Teaching Unit", args["teaching_unit"], "semestre")


def _get_config():
    setting = frappe.get_single("Udshed Setting")
    return {
        "methode_calcul_cc": setting.methode_calcul_cc or "Moyenne arithmétique",
        "nb_meilleures_notes_cc": setting.nb_meilleures_notes_cc or 2,
        "exiger_programmation": bool(setting.exiger_programmation),
        "note_max": 20,
        "seuil_licence": setting.seuil_validation_licence or 50,
        "seuil_master": setting.seuil_validation_master or 60,
    }


def _get_niveau_name(filiere, niveau_label):
    """Retourne le name (ID) d'un niveau depuis son label, pour un niveau donné."""
    if not niveau_label:
        return None
    filiere_doc = frappe.get_doc("Field of study", filiere)
    for row in filiere_doc.field_of_study_level:
        if row.level == niveau_label:
            return row.name
    return None


def _get_calendar_defaut():
    calendar = frappe.db.get_value("Calendar Planing", {}, "name")
    if calendar:
        return calendar
    doc = frappe.new_doc("Calendar Planing")
    doc.nom_du_planing = "Défaut"
    doc.insert(ignore_permissions=True)
    return doc.name


def _get_or_create_session(args, type_dexamen):
    """Retrouve ou crée la Session Examen d'un type donné pour (année, semestre, filière, niveau)."""
    academic_year = args["academic_year"]
    semestre = args.get("semestre")
    filiere = args["filiere"]
    niveau = _get_niveau_name(filiere, args["niveau"])
    if not niveau:
        frappe.throw(_("Le niveau {0} est introuvable pour la filière {1}.").format(args["niveau"], filiere))

    sessions = frappe.get_all(
        "Session Examen",
        filters={"academic_year": academic_year, "semestre": semestre, "type_dexamen": type_dexamen},
        fields=["name"],
    )

    for session in sessions:
        if frappe.db.exists(
            "Session Examen Field of study Level",
            {"parent": session.name, "filiere": filiere, "niveau": niveau},
        ):
            return session.name

    if sessions:
        # Aucune session ne couvre encore cette filière/niveau : on l'ajoute à la première
        doc = frappe.get_doc("Session Examen", sessions[0].name)
        doc.append("classes_concernees", {"filiere": filiere, "niveau": niveau})
        doc.save(ignore_permissions=True)
        return doc.name

    doc = frappe.new_doc("Session Examen")
    doc.academic_year = academic_year
    doc.calendar = _get_calendar_defaut()
    doc.type_dexamen = type_dexamen
    doc.semestre = semestre
    doc.date_debut = today()
    doc.date_de_fin = today()
    doc.append("classes_concernees", {"filiere": filiere, "niveau": niveau})
    doc.insert(ignore_permissions=True)
    return doc.name


def _chercher_session(args, type_dexamen):
    """Retrouve (sans la créer) la Session Examen d'un type donné pour la filière/niveau."""
    academic_year = args["academic_year"]
    semestre = args.get("semestre")
    filiere = args["filiere"]
    niveau = _get_niveau_name(filiere, args["niveau"])
    sessions = frappe.get_all(
        "Session Examen",
        filters={"academic_year": academic_year, "semestre": semestre, "type_dexamen": type_dexamen},
        fields=["name"],
    )
    for session in sessions:
        if frappe.db.exists(
            "Session Examen Field of study Level",
            {"parent": session.name, "filiere": filiere, "niveau": niveau},
        ):
            return session.name
    return sessions[0]["name"] if sessions else None


def _verifier_programmation(academic_year, teaching_unit):
    from frappe.query_builder import DocType

    planning = DocType("Planning Item")
    query = (
        frappe.qb.from_(planning)
        .select(planning.name)
        .where(
            (planning.cours == teaching_unit)
            & (planning.academic_year == academic_year)
            & (planning.type.isin(TYPES_EVAL))
        )
        .limit(1)
    )
    return bool(query.run())


def _verifier_programmation_requise(academic_year, teaching_unit):
    """Lève une erreur si la configuration exige que l'évaluation soit programmée."""
    config = _get_config()
    if config["exiger_programmation"] and not _verifier_programmation(academic_year, teaching_unit):
        frappe.throw(
            _(
                "L'UE {0} n'a aucune évaluation (CC, examen ou rattrapage) programmée pour {1} "
                "dans le planning académique."
            ).format(teaching_unit, academic_year)
        )


def _verifier_acces_enseignant(teaching_unit):
    """Vérifie que l'utilisateur connecté a le droit de gérer les notes d'une UE.

    - Administrateur / System Manager : accès total
    - Coordonnateur / Gestionnaire de planning : accès (gestion de filière)
    - Enseignant : uniquement les UE qu'il dispense (Course Teacher Item)
    """
    user = frappe.session.user
    roles = frappe.get_roles(user)
    if user == "Administrator" or "System Manager" in roles:
        return
    if (
        "Planning Manager" in roles
        or "Coordonateur" in roles
        or "Coordinateur" in roles
    ):
        return
    if "Teacher" in roles:
        enseignant = _get_enseignant_courant()
        if enseignant and frappe.db.exists(
            "Course Teacher Item",
            {"parent": teaching_unit, "enseignant": enseignant},
        ):
            return
        frappe.throw(
            _(
                "Vous n'êtes pas autorisé(e) à gérer les notes de l'UE <b>{0}</b> "
                "(cet enseignant ne dispense pas ce cours)."
            ).format(teaching_unit)
        )
    frappe.throw(frappe.PermissionError)


def _valider_lignes(rows):
    """Valide les lignes avant enregistrement : note dans [0,20], aucune ligne en double.

    Retourne la liste des lignes normalisées (accepte aussi une chaîne JSON).
    """
    if isinstance(rows, str):
        try:
            rows = frappe.parse_json(rows) or []
        except Exception:
            frappe.throw(_("Les lignes de notes envoyées ne sont pas valides (format JSON)."))
    if not isinstance(rows, list):
        frappe.throw(_("Les lignes de notes doivent être une liste."))
    vus = set()
    for row in rows:
        student = row.get("student")
        if not student:
            frappe.throw(_("Une ligne d'étudiant est vide : impossible d'enregistrer."))
        if student in vus:
            frappe.throw(
                _("L'étudiant <b>{0}</b> apparaît plusieurs fois dans les notes à enregistrer.").format(student)
            )
        vus.add(student)

        for item in row.get("notes_cc") or []:
            valeur = item.get("note_cc")
            if valeur in (None, ""):
                continue
            valeur = flt(valeur)
            if valeur < 0 or valeur > NOTE_MAX:
                frappe.throw(
                    _("La note CC <b>{0}</b> de l'étudiant <b>{1}</b> doit être comprise entre 0 et {2}.").format(
                        item.get("cc_label") or "CC", student, NOTE_MAX
                    )
                )

        for champ in ("note_examen", "note_tp", "note_examen_rattrapage", "note_cctp", "note_examtp", "cc"):
            valeur = row.get(champ)
            if valeur in (None, ""):
                continue
            valeur = flt(valeur)
            if valeur < 0 or valeur > NOTE_MAX:
                frappe.throw(
                    _("La note <b>{0}</b> de l'étudiant <b>{1}</b> doit être comprise entre 0 et {2}.").format(
                        champ.replace("_", " "), student, NOTE_MAX
                    )
                )
    return rows


def _credits_ue(teaching_unit, filiere=None, niveau=None):
    """Crédits d'une UE : priorité à la grille de matière (course_poid filière/niveau), sinon le champ Teaching Unit."""
    if filiere and niveau:
        fos_doc = frappe.get_doc("Field of study", filiere)
        niveau_name = None
        for row in fos_doc.field_of_study_level or []:
            if row.level == niveau:
                niveau_name = row.name
                break
        if niveau_name:
            poid = frappe.db.get_value(
                "Course Field of study level item",
                {"parent": teaching_unit, "filiere": filiere, "niveau": niveau_name},
                "course_poid",
            )
            if poid:
                return poid
    credits = frappe.db.get_value("Teaching Unit", teaching_unit, "credits")
    return credits or 0


def _get_ue_info(teaching_unit, filiere=None, niveau=None):
    tu = frappe.get_doc("Teaching Unit", teaching_unit)
    course = {}
    if tu.course:
        course_doc = frappe.get_value("Course", tu.course, ["code", "intitule"], as_dict=True)
        course = course_doc or {}

    enseignants = []
    vus = set()
    for row in tu.table_enseignant or []:
        if not row.enseignant or row.enseignant in vus:
            continue
        vus.add(row.enseignant)
        nom = frappe.db.get_value("Teacher", row.enseignant, "full_name") or ""
        enseignants.append({"name": row.enseignant, "full_name": nom})

    return {
        "name": tu.name,
        "code": course.get("code") or "",
        "intitule": course.get("intitule") or tu.intitule_cours or tu.name,
        "credits": _credits_ue(teaching_unit, filiere, niveau),
        "type_ue": tu.type_ue,
        "semestre": tu.semestre,
        "enseignants": enseignants,
    }


def _get_etudiants(academic_year, filiere, niveau_label, teaching_unit):
    """Étudiants réellement inscrits au cours (teaching_unit) dans ce contexte académique.

    La liste est construite depuis les réinscriptions académiques validées de la
    filière/niveau, en ne conservant que les étudiants dont la fiche de réinscription
    contient ce cours avec un statut « Inscrit » (ni « Dispensé », ni « Reporté »).
    """
    regs = frappe.get_all(
        "Academic Reregistration",
        filters={
            "academic_year": academic_year,
            "filiere": filiere,
            "niveau": niveau_label,
            "statut": "Validée",
        },
        fields=["name", "student"],
    )
    if not regs:
        return []

    parents = [r.name for r in regs]
    inscrits = frappe.get_all(
        "Reregistration Course Item",
        filters={
            "parent": ["in", parents],
            "teaching_unit": teaching_unit,
            "statut": ["not in", ["Dispensé", "Reporté"]],
        },
        fields=["parent"],
    )
    inscrits_parents = {i.parent for i in inscrits}

    students = [r.student for r in regs if r.name in inscrits_parents]

    resultats = []
    for student in students:
        info = frappe.db.get_value("Student", student, ["matricule", "nom", "prenom"], as_dict=True)
        if not info:
            continue
        resultats.append(
            {
                "student": student,
                "matricule": info.matricule or student,
                "nom": info.nom or "",
                "prenom": info.prenom or "",
            }
        )

    resultats.sort(key=lambda s: (s["matricule"] or s["student"]).lower())
    return resultats


def _get_or_create_note(session, student, args):
    teaching_unit = args["teaching_unit"]
    existing = frappe.db.get_value(
        "Session Examen Note",
        {"session_examen": session, "student": student, "teaching_unit": teaching_unit},
        "name",
    )
    if existing:
        return frappe.get_doc("Session Examen Note", existing)

    doc = frappe.new_doc("Session Examen Note")
    doc.session_examen = session
    doc.student = student
    doc.teaching_unit = teaching_unit
    doc.filiere = args["filiere"]
    doc.niveau = _get_niveau_name(args["filiere"], args["niveau"])
    doc.type_ue = frappe.db.get_value("Teaching Unit", teaching_unit, "type_ue")
    doc.insert(ignore_permissions=True)
    return doc


def _charger_notes(students, teaching_unit, sessions):
    resultats = {"CC": {}, "Examen": {}, "TP": {}, "Rattrapage": {}}
    if not students:
        return resultats

    student_names = [s["student"] for s in students]
    session_names = list(sessions.values())
    notes = frappe.get_all(
        "Session Examen Note",
        filters={
            "session_examen": ["in", session_names],
            "student": ["in", student_names],
            "teaching_unit": teaching_unit,
        },
        fields=[
            "name",
            "session_examen",
            "student",
            "note_cc_moyenne",
            "cc_saisi",
            "note_examen",
            "examen_saisi",
            "note_tp",
            "tp_saisi",
            "note_cctp",
            "cctp_saisi",
            "note_examtp",
            "examtp_saisi",
            "note_examen_rattrapage",
            "rattrapage_saisi",
            "date_rattrapage",
            "note_examen_active",
            "note_finale",
            "note_pct",
            "grade",
            "point",
            "mention",
            "statut",
        ],
    )
    if not notes:
        return resultats

    note_names = [n.name for n in notes]
    items_cc = frappe.get_all(
        "Note CC Item",
        filters={"parent": ["in", note_names]},
        fields=["parent", "cc_label", "cc_weight", "note_cc"],
        order_by="idx asc",
    )
    items_by_note = {}
    for item in items_cc:
        items_by_note.setdefault(item.parent, []).append(
            {
                "cc_label": item.cc_label,
                "cc_weight": item.cc_weight,
                "note_cc": item.note_cc,
            }
        )

    par_session = {}
    for n in notes:
        par_session.setdefault(n.session_examen, {})[n.student] = n

    for s in students:
        nom = s["student"]

        cc_note = (par_session.get(sessions["cc"]) or {}).get(nom)
        if cc_note:
            resultats["CC"][nom] = {
                "notes_cc": items_by_note.get(cc_note.name, []),
                "note_cc_moyenne": cc_note.note_cc_moyenne,
                "cc_saisi": cc_note.cc_saisi,
                "note_finale": cc_note.note_finale,
                "grade": cc_note.grade,
                "point": cc_note.point,
                "mention": cc_note.mention,
            }

        normale = (par_session.get(sessions["normale"]) or {}).get(nom)
        if normale:
            resultats["Examen"][nom] = {
                "note_cc_moyenne": normale.note_cc_moyenne,
                "cc_saisi": normale.cc_saisi,
                "note_examen": normale.note_examen,
                "examen_saisi": normale.examen_saisi,
                "note_tp": normale.note_tp,
                "tp_saisi": normale.tp_saisi,
                "note_cctp": normale.note_cctp,
                "cctp_saisi": normale.cctp_saisi,
                "note_examtp": normale.note_examtp,
                "examtp_saisi": normale.examtp_saisi,
                "note_examen_active": normale.note_examen_active,
                "note_finale": normale.note_finale,
                "note_pct": normale.note_pct,
                "grade": normale.grade,
                "point": normale.point,
                "mention": normale.mention,
                "statut": normale.statut,
            }
            resultats["TP"][nom] = {
                "note_tp": normale.note_tp,
                "tp_saisi": normale.tp_saisi,
            }

        rattrapage = (par_session.get(sessions["rattrapage"]) or {}).get(nom)
        initiale = normale.note_examen if normale and normale.examen_saisi else None
        if rattrapage:
            resultats["Rattrapage"][nom] = {
                "note_examen": rattrapage.note_examen
                if rattrapage.examen_saisi
                else initiale,
                "note_examen_rattrapage": rattrapage.note_examen_rattrapage,
                "rattrapage_saisi": rattrapage.rattrapage_saisi,
                "date_rattrapage": rattrapage.date_rattrapage,
                "note_examen_active": rattrapage.note_examen_active,
                "note_finale": rattrapage.note_finale,
                "note_pct": rattrapage.note_pct,
                "grade": rattrapage.grade,
                "point": rattrapage.point,
                "mention": rattrapage.mention,
            }
        elif initiale is not None:
            resultats["Rattrapage"][nom] = {
                "note_examen": initiale,
                "note_examen_rattrapage": None,
                "rattrapage_saisi": 0,
                "date_rattrapage": None,
                "note_examen_active": initiale,
                "note_finale": None,
                "note_pct": None,
                "grade": None,
                "point": None,
                "mention": None,
            }

    return resultats


def _verifier_non_publiee(session):
    """Publication non verrouillante.

    Une session publiée reste modifiable : des requêtes (réclamations) et des
    corrections de notes peuvent survenir après la publication. Ce contrôle ne
    bloque donc plus la saisie, quelle que soit l'état de la session.
    """


def _passer_saisi(note):
    """Fait passer une note au statut « Saisi » dès qu'une saisie est enregistrée.

    Le statut ne recule jamais : une note « Validé » ou « Publié » le reste.
    """
    if note.statut in (None, "", "Brouillon"):
        note.statut = "Saisi"


def _verifier_cc_modifiable(note):
    """Le CC reste modifiable, même après validation/publication (requêtes)."""


def _copier_cc_dans_note(args, student, note):
    """Reproduit le CC (lecture seule) de la session CC sur la note cible.

    Le CC est saisi dans sa propre session (Contrôle Continu) ; pour que la
    note de matière de la session normale soit complète (CC + Examen + TP),
    ses notes de CC sont synchronisées depuis la session CC. La session CC
    reste la seule source : on ne modifie jamais la note source. La
    synchronisation est systématique (elle remplace les notes de CC déjà
    copiées) afin de refléter les dernières valeurs saisies.
    """
    session_cc = _chercher_session(args, TYPE_CC)
    if not session_cc:
        return
    source = frappe.db.get_value(
        "Session Examen Note",
        {
            "session_examen": session_cc,
            "student": student,
            "teaching_unit": args["teaching_unit"],
        },
        "name",
    )
    if not source:
        return
    source_doc = frappe.get_doc("Session Examen Note", source)
    note.set("notes_cc", [])
    for cc in source_doc.notes_cc:
        if cc.note_cc is not None:
            note.append(
                "notes_cc",
                {
                    "cc_label": cc.cc_label,
                    "cc_weight": cc.cc_weight,
                    "note_cc": cc.note_cc,
                },
            )


def _sauvegarder_cc(args, rows):
    rows = _valider_lignes(rows)
    _verifier_programmation_requise(args["academic_year"], args["teaching_unit"])
    session = _get_or_create_session(args, TYPE_CC)
    _verifier_non_publiee(session)
    for row in rows:
        note = _get_or_create_note(session, row["student"], args)
        _verifier_cc_modifiable(note)
        note.set("notes_cc", [])
        labels_vus = set()
        for item in row.get("notes_cc") or []:
            valeur = item.get("note_cc")
            if valeur in (None, ""):
                continue
            label = item.get("cc_label") or "CC"
            if label in labels_vus:
                continue
            labels_vus.add(label)
            # Les poids sont définis par la Grade Formula : la moyenne CC et la
            # pondération de la composante sont calculées par le moteur central.
            # Le poids par note CC est donc toujours 1 (moyenne arithmétique),
            # jamais fourni par l'enseignant.
            note.append(
                "notes_cc",
                {
                    "cc_label": label,
                    "cc_weight": 1,
                    "note_cc": flt(valeur),
                },
            )
        notes_saisies = [i for i in (row.get("notes_cc") or []) if i.get("note_cc") not in (None, "")]
        note.cc_saisi = 1 if notes_saisies else 0
        _passer_saisi(note)
        note.save(ignore_permissions=True)
    return len(rows)


def _sauvegarder_examen(args, rows):
    rows = _valider_lignes(rows)
    _verifier_programmation_requise(args["academic_year"], args["teaching_unit"])
    session = _get_or_create_session(args, TYPE_NORMALE)
    _verifier_non_publiee(session)
    for row in rows:
        note = _get_or_create_note(session, row["student"], args)
        _copier_cc_dans_note(args, row["student"], note)
        note.note_examen = row.get("note_examen")
        note.examen_saisi = 1 if row.get("note_examen") not in (None, "") else 0
        if "note_tp" in row:
            note.note_tp = row.get("note_tp")
            note.tp_saisi = 1 if row.get("note_tp") not in (None, "") else 0
        _passer_saisi(note)
        note.save(ignore_permissions=True)
    return len(rows)


def _sauvegarder_tp(args, rows):
    rows = _valider_lignes(rows)
    _verifier_programmation_requise(args["academic_year"], args["teaching_unit"])
    session = _get_or_create_session(args, TYPE_NORMALE)
    _verifier_non_publiee(session)
    for row in rows:
        note = _get_or_create_note(session, row["student"], args)
        _copier_cc_dans_note(args, row["student"], note)
        note.note_tp = row.get("note_tp")
        note.tp_saisi = 1 if row.get("note_tp") not in (None, "") else 0
        _passer_saisi(note)
        note.save(ignore_permissions=True)
    return len(rows)


def _sauvegarder_rattrapage(args, rows):
    rows = _valider_lignes(rows)
    _verifier_programmation_requise(args["academic_year"], args["teaching_unit"])
    session = _get_or_create_session(args, TYPE_RATTRAPAGE)
    session_normale = _get_or_create_session(args, TYPE_NORMALE)
    _verifier_non_publiee(session)
    for row in rows:
        note = _get_or_create_note(session, row["student"], args)
        note.note_examen_rattrapage = row.get("note_examen_rattrapage")
        note.rattrapage_saisi = 1 if row.get("note_examen_rattrapage") not in (None, "") else 0

        # Conserver la note d'examen initiale sur la note de rattrapage afin que
        # la note retenue (note_examen_active = max(initiale, rattrapage)) soit correcte.
        initiale = frappe.db.get_value(
            "Session Examen Note",
            {
                "session_examen": session_normale,
                "student": row["student"],
                "teaching_unit": args["teaching_unit"],
            },
            "name",
        )
        if initiale:
            normale_doc = frappe.get_doc("Session Examen Note", initiale)
            note.note_examen = normale_doc.note_examen
            if normale_doc.note_examen is not None:
                note.examen_saisi = 1

            # Les composantes acquises en session normale sont conservées :
            # TP, Rapport et Compétence ne sont pas à ressaisir au rattrapage.
            # Seul un drapeau *saisi (ou une valeur non nulle) compte : le
            # stockage Currency ramène les champs non saisis à 0.0.
            note.note_tp = normale_doc.note_tp
            if normale_doc.tp_saisi:
                note.tp_saisi = 1
            if normale_doc.note_rapport:
                note.note_rapport = normale_doc.note_rapport
            if normale_doc.note_competence:
                note.note_competence = normale_doc.note_competence

        # Le CC provient toujours de la session CC (source unique) ; à défaut
        # d'une session CC dédiée, on reprend le CC de la note de la session
        # normale afin que la moyenne CC soit conservée au rattrapage.
        _copier_cc_dans_note(args, row["student"], note)
        if not note.get("notes_cc") and initiale:
            for cc in normale_doc.notes_cc:
                if cc.note_cc is not None:
                    note.append(
                        "notes_cc",
                        {
                            "cc_label": cc.cc_label,
                            "cc_weight": cc.cc_weight,
                            "note_cc": cc.note_cc,
                        },
                    )

        if row.get("date_rattrapage"):
            note.date_rattrapage = row.get("date_rattrapage")
        _passer_saisi(note)
        note.save(ignore_permissions=True)
    return len(rows)


@frappe.whitelist()
def charger_data(academic_year, filiere, niveau, semestre, teaching_unit):
    """Charge la configuration, l'UE, les étudiants et toutes les notes pour la page de saisie."""
    _verifier_acces_enseignant(teaching_unit)
    config = _get_config()

    if config["exiger_programmation"] and not _verifier_programmation(academic_year, teaching_unit):
        frappe.throw(
            _(
                "L'UE {0} n'a aucune évaluation (CC, examen ou rattrapage) programmée pour {1} "
                "dans le planning académique."
            ).format(teaching_unit, academic_year)
        )

    ue_info = _get_ue_info(teaching_unit, filiere, niveau)
    students = _get_etudiants(academic_year, filiere, niveau, teaching_unit)

    args = {
        "academic_year": academic_year,
        "semestre": _semestre_effectif({"academic_year": academic_year, "semestre": semestre, "teaching_unit": teaching_unit}),
        "filiere": filiere,
        "niveau": niveau,
        "teaching_unit": teaching_unit,
    }
    semestre_effectif = args["semestre"]

    sessions = {
        "cc": _get_or_create_session(args, TYPE_CC),
        "normale": _get_or_create_session(args, TYPE_NORMALE),
        "rattrapage": _get_or_create_session(args, TYPE_RATTRAPAGE),
    }

    sessions_meta = {}
    for cle, session in sessions.items():
        sessions_meta[cle] = {
            "name": session,
            "statut": frappe.db.get_value("Session Examen", session, "statut"),
        }

    notes = _charger_notes(students, teaching_unit, sessions)

    return {
        "config": config,
        "ue_info": ue_info,
        "formule": _formule_detectee(teaching_unit, filiere, niveau, notes, students),
        "students": students,
        "sessions": sessions_meta,
        "notes": notes,
        "lignes": _lignes_unifiees(students, notes),
    }


def _lignes_unifiees(students, notes):
    """Lignes de la saisie unifiée : une ligne par étudiant inscrit.

    La source unique est la note de la session normale (CC, CCTP, EXAMTP,
    EXAM) ; à défaut de CC saisi directement sur la session normale, on
    retombe sur la note de la session CC (données historiques). Les valeurs
    ne sont présentées que si le drapeau *saisi est posé : le stockage
    Currency ramène les champs non saisis à 0.0.
    """
    notes_cc = (notes or {}).get("CC", {})
    lignes = []
    for s in students or []:
        nom = s["student"]
        ex = dict((notes or {}).get("Examen", {}).get(nom, {}))
        cc = notes_cc.get(nom, {})
        if not ex.get("cc_saisi") and cc.get("note_cc_moyenne"):
            ex["note_cc_moyenne"] = cc.get("note_cc_moyenne")
            ex["cc_saisi"] = 1
        lignes.append(
            {
                "matricule": s.get("matricule") or "",
                "nom": s.get("nom") or "",
                "prenom": s.get("prenom") or "",
                "student": nom,
                "cc": ex.get("note_cc_moyenne") if ex.get("cc_saisi") else None,
                "cctp": ex.get("note_cctp") if ex.get("cctp_saisi") else None,
                "examtp": ex.get("note_examtp") if ex.get("examtp_saisi") else None,
                "examen": ex.get("note_examen") if ex.get("examen_saisi") else None,
                "note_finale": ex.get("note_finale"),
                "note_pct": ex.get("note_pct"),
                "grade": ex.get("grade"),
                "point": ex.get("point"),
                "mention": ex.get("mention"),
                "statut": ex.get("statut"),
            }
        )
    return lignes


@frappe.whitelist()
def get_ues(academic_year, filiere, niveau, semestre=None, teacher=None):
    """Retourne les UE d'un (année, filière, niveau, semestre) pour les filtres de la page."""
    from frappe.query_builder import DocType

    filiere_doc = frappe.get_doc("Field of study", filiere)
    niveau_names = [row.name for row in filiere_doc.field_of_study_level if row.level == niveau]
    if not niveau_names:
        return []

    TeachingUnit = DocType("Teaching Unit")
    CourseLevel = DocType("Course Field of study level item")
    query = (
        frappe.qb.from_(TeachingUnit)
        .join(CourseLevel)
        .on(CourseLevel.parent == TeachingUnit.name)
        .select(
            TeachingUnit.name,
            TeachingUnit.intitule_cours,
            TeachingUnit.course,
            TeachingUnit.semestre,
            TeachingUnit.credits,
        )
        .where(
            (TeachingUnit.academic_year == academic_year)
            & (CourseLevel.filiere == filiere)
            & (CourseLevel.niveau.isin(niveau_names))
        )
        .distinct()
    )
    if semestre:
        query = query.where(TeachingUnit.semestre == semestre)
    if teacher:
        CourseTeacherItem = DocType("Course Teacher Item")
        query = (
            query.join(CourseTeacherItem)
            .on(CourseTeacherItem.parent == TeachingUnit.name)
            .where(CourseTeacherItem.enseignant == teacher)
        )

    resultats = []
    for ue in query.run(as_dict=True):
        code = ""
        if ue.course:
            code = frappe.db.get_value("Course", ue.course, "code") or ""
        resultats.append(
            {
                "name": ue.name,
                "intitule": ue.intitule_cours or ue.name,
                "code": code,
                "semestre": ue.semestre,
                "credits": ue.credits,
            }
        )
    resultats.sort(key=lambda u: (u["intitule"] or u["name"]).lower())
    return resultats


@frappe.whitelist()
def get_enseignant_courant():
    """Infos sur l'enseignant lié à l'utilisateur connecté (filtre Enseignant / droits)."""
    enseignant = _get_enseignant_courant()
    if not enseignant:
        return None
    return {"name": enseignant, "full_name": _get_enseignant_full_name(enseignant)}


@frappe.whitelist()
def get_enseignants(academic_year, filiere, niveau, semestre=None):
    """Liste les enseignants qui dispensent une UE dans le contexte sélectionné."""
    ues = get_ues(academic_year, filiere, niveau, semestre)
    if not ues:
        return []

    noms = [u["name"] for u in ues]
    rows = frappe.get_all(
        "Course Teacher Item",
        filters={"parent": ["in", noms]},
        fields=["parent", "enseignant"],
        distinct=True,
    )
    par_enseignant = {}
    for row in rows:
        if not row.enseignant:
            continue
        par_enseignant.setdefault(row.enseignant, set()).add(row.parent)

    resultats = [
        {"name": e, "full_name": _get_enseignant_full_name(e), "nb_ues": len(ues_p)}
        for e, ues_p in par_enseignant.items()
    ]
    resultats.sort(key=lambda t: (t["full_name"] or t["name"]).lower())
    return resultats


@frappe.whitelist()
def calculer_apercu(student, teaching_unit, notes_cc=None, note_examen=None, note_tp=None, note_examen_rattrapage=None, note_cctp=None, note_examtp=None):
    """Calcule à la volée (sans enregistrer) moyenne CC, note retenue, note finale, % et grade.

    La combinaison des évaluations renseignées est détectée automatiquement et
    la formule est résolue dans la Grade Formula (poids de la configuration).
    Utilise le même moteur central que l'enregistrement (Session Examen Note)
    afin que l'aperçu affiché au professeur soit strictement identique au
    résultat qui sera réellement enregistré.
    """
    _verifier_acces_enseignant(teaching_unit)
    from udshed.udshed.doctype.session_examen_note.session_examen_note import SessionExamenNote  # noqa: F401

    type_ue = frappe.db.get_value("Teaching Unit", teaching_unit, "type_ue") or "Sans TP"

    doc = frappe.new_doc("Session Examen Note")
    doc.student = student
    doc.teaching_unit = teaching_unit
    doc.type_ue = type_ue

    if isinstance(notes_cc, str):
        try:
            notes_cc = frappe.parse_json(notes_cc) or []
        except Exception:
            notes_cc = []
    for item in notes_cc or []:
        valeur = item.get("note_cc")
        if valeur in (None, ""):
            continue
        doc.append(
            "notes_cc",
            {
                "cc_label": item.get("cc_label") or "CC",
                "cc_weight": flt(item.get("cc_weight") or 1),
                "note_cc": flt(valeur),
            },
        )

    if note_examen not in (None, ""):
        doc.note_examen = flt(note_examen)
    if note_tp not in (None, ""):
        doc.note_tp = flt(note_tp)
    if note_cctp not in (None, ""):
        doc.note_cctp = flt(note_cctp)
    if note_examtp not in (None, ""):
        doc.note_examtp = flt(note_examtp)
    if note_examen_rattrapage not in (None, ""):
        doc.note_examen_rattrapage = flt(note_examen_rattrapage)

    doc.deriver_drapeaux_saisie()
    doc.calculer_note_cc_moyenne()
    doc.calculer_note_examen_active()
    doc.calculer_note_finale()
    doc.determiner_grade()

    from udshed.grade_calculation import get_student_cycle

    formula = doc._formule_note()
    cycle = get_student_cycle(student)
    formule = (
        _serialiser_formule(formula, cycle, doc.combinaison_detectee())
        if formula
        else None
    )

    return {
        "note_cc_moyenne": doc.note_cc_moyenne,
        "note_examen_active": doc.note_examen_active,
        "note_finale": doc.note_finale,
        "note_pct": doc.note_pct,
        "grade": doc.grade,
        "point": doc.point,
        "mention": doc.mention,
        "type_resultat": doc.type_resultat,
        "combinaison": doc.combinaison_detectee(),
        "formule_manquante": bool(doc.formule_manquante()),
        "formule": formule,
    }


@frappe.whitelist()
def enregistrer_cc(academic_year, filiere, niveau, semestre, teaching_unit, rows):
    """Enregistre les notes de contrôle continu pour tous les étudiants."""
    _verifier_acces_enseignant(teaching_unit)
    args = {"academic_year": academic_year, "filiere": filiere, "niveau": niveau, "semestre": _semestre_effectif({"academic_year": academic_year, "semestre": semestre, "teaching_unit": teaching_unit}), "teaching_unit": teaching_unit}
    n = _sauvegarder_cc(args, rows)
    return {"saved": n}


@frappe.whitelist()
def enregistrer_evaluations(academic_year, filiere, niveau, semestre, teaching_unit, rows):
    """Enregistre la saisie unifiée (CC, CCTP, EXAMTP, EXAM) d'une UE.

    Chaque ligne : {"student", "cc", "note_cctp", "note_examtp", "note_examen"}.
    Les quatre évaluations sont stockées sur la note de la session normale : le
    CC comme une note CC unique (poids 1), les autres dans leurs champs dédiés.
    La combinaison est détectée et la formule résolue par le moteur central.
    """
    _verifier_acces_enseignant(teaching_unit)
    args = {"academic_year": academic_year, "filiere": filiere, "niveau": niveau, "semestre": _semestre_effectif({"academic_year": academic_year, "semestre": semestre, "teaching_unit": teaching_unit}), "teaching_unit": teaching_unit}
    rows = _valider_lignes(rows)
    _verifier_programmation_requise(args["academic_year"], args["teaching_unit"])
    session = _get_or_create_session(args, TYPE_NORMALE)
    _verifier_non_publiee(session)
    n = 0
    for row in rows:
        note = _get_or_create_note(session, row["student"], args)
        _copier_cc_dans_note(args, row["student"], note)
        cc = row.get("cc")
        if cc not in (None, ""):
            note.set("notes_cc", [])
            note.append("notes_cc", {"cc_label": "CC", "cc_weight": 1, "note_cc": flt(cc)})
        for champ in ("note_cctp", "note_examtp", "note_examen"):
            valeur = row.get(champ)
            setattr(note, champ, None if valeur in (None, "") else flt(valeur))
        note.cctp_saisi = 1 if row.get("note_cctp") not in (None, "") else 0
        note.examtp_saisi = 1 if row.get("note_examtp") not in (None, "") else 0
        note.examen_saisi = 1 if row.get("note_examen") not in (None, "") else 0
        _passer_saisi(note)
        note.save(ignore_permissions=True)
        n += 1
    return {"saved": n}


@frappe.whitelist()
def enregistrer_examen(academic_year, filiere, niveau, semestre, teaching_unit, rows):
    """Enregistre les notes de l'examen de session normale."""
    _verifier_acces_enseignant(teaching_unit)
    args = {"academic_year": academic_year, "filiere": filiere, "niveau": niveau, "semestre": _semestre_effectif({"academic_year": academic_year, "semestre": semestre, "teaching_unit": teaching_unit}), "teaching_unit": teaching_unit}
    n = _sauvegarder_examen(args, rows)
    return {"saved": n}


@frappe.whitelist()
def enregistrer_tp(academic_year, filiere, niveau, semestre, teaching_unit, rows):
    """Enregistre les notes de travaux pratiques."""
    _verifier_acces_enseignant(teaching_unit)
    args = {"academic_year": academic_year, "filiere": filiere, "niveau": niveau, "semestre": _semestre_effectif({"academic_year": academic_year, "semestre": semestre, "teaching_unit": teaching_unit}), "teaching_unit": teaching_unit}
    n = _sauvegarder_tp(args, rows)
    return {"saved": n}


@frappe.whitelist()
def enregistrer_rattrapage(academic_year, filiere, niveau, semestre, teaching_unit, rows):
    """Enregistre les notes de rattrapage."""
    _verifier_acces_enseignant(teaching_unit)
    args = {"academic_year": academic_year, "filiere": filiere, "niveau": niveau, "semestre": _semestre_effectif({"academic_year": academic_year, "semestre": semestre, "teaching_unit": teaching_unit}), "teaching_unit": teaching_unit}
    n = _sauvegarder_rattrapage(args, rows)
    return {"saved": n}


@frappe.whitelist()
def valider_notes(session):
    """Valide toutes les notes d'une session d'examen (statut -> « Validé »).

    Étape obligatoire avant la publication : une session ne peut être
    publiée que si toutes ses notes sont validées.

    Args:
        session: Nom du document Session Examen

    Returns:
        dict: {"session": ..., "validated": n}
    """
    names = frappe.get_all(
        "Session Examen Note",
        filters={"session_examen": session, "statut": ["in", ["Brouillon", "Saisi"]]},
        pluck="name",
    )
    for name in names:
        frappe.db.set_value("Session Examen Note", name, "statut", "Validé")

    return {"session": session, "validated": len(names)}


@frappe.whitelist()
def publier_session(session):
    """Publie une session d'examen.

    Toutes les notes de la session doivent d'abord être validées
    (statut « Validé »). Elles passent alors au statut « Publié » et
    deviennent consultables au babillard public. La publication n'est pas
    définitive : les notes restent modifiables (corrections / réclamations)
    après publication.

    Args:
        session: Nom du document Session Examen

    Returns:
        dict: {"name", "statut", "date_publication"}
    """
    doc = frappe.get_doc("Session Examen", session)
    if doc.statut == "Publiée":
        return {"name": doc.name, "statut": doc.statut, "date_publication": doc.date_publication}

    non_validees = frappe.db.count(
        "Session Examen Note",
        {"session_examen": session, "statut": ["not in", ["Validé", "Publié"]]},
    )
    if non_validees:
        frappe.throw(
            _("Impossible de publier : {0} note(s) de la session {1} ne sont pas validées. "
              "Validez d'abord les notes avant de publier.").format(non_validees, session)
        )

    frappe.db.set_value(
        "Session Examen Note",
        {"session_examen": session, "statut": "Validé"},
        "statut",
        "Publié",
    )
    doc.db_set("statut", "Publiée")
    doc.db_set("date_publication", today())
    return {"name": doc.name, "statut": doc.statut, "date_publication": doc.date_publication}


def _colonnes_cc(cc_columns):
    """Normalise la liste des colonnes CC envoyee par le client (labels [+ poids])."""
    if isinstance(cc_columns, str):
        try:
            cc_columns = frappe.parse_json(cc_columns)
        except Exception:
            cc_columns = None
    return [
        c if isinstance(c, dict) else {"label": c, "weight": 1}
        for c in (cc_columns or [])
    ]


def _en_tetes(type_dexamen, cc_columns):
    colonnes = [_("Matricule"), _("Nom"), _("Prénom"), _("Crédit")]
    if type_dexamen == "CC":
        return colonnes + [c.get("label") for c in cc_columns] + [_("Moyenne CC")]
    if type_dexamen == "Examen":
        return colonnes + [_("Note d'examen")]
    if type_dexamen == "TP":
        return colonnes + [_("Note de TP")]
    if type_dexamen == "Rattrapage":
        return [_("Matricule"), _("Nom"), _("Prénom"), _("Note d'examen initiale"), _("Note de rattrapage"), _("Note retenue")]
    return colonnes


def _ue_a_tp(teaching_unit):
    """Vrai si le cours (Teaching Unit) a des Travaux Pratiques (type UE « Avec TP »)."""
    if not teaching_unit:
        return False
    return (frappe.db.get_value("Teaching Unit", teaching_unit, "type_ue") or "") == "Avec TP"


def _entetes_avec_tp(entetes, type_dexamen, teaching_unit):
    """Insère la colonne « Note de TP » dans un modèle d'examen si le cours a un TP."""
    if type_dexamen != "Examen" or not _ue_a_tp(teaching_unit):
        return entetes
    entetes = list(entetes)
    if _("Note de TP") not in entetes:
        pos = entetes.index(_("Note d'examen")) if _("Note d'examen") in entetes else len(entetes)
        entetes.insert(pos, _("Note de TP"))
    return entetes


def _repondre_xlsx(rows, filename):
    from frappe.utils.xlsxutils import make_xlsx

    xlsx = make_xlsx(rows, "Notes")
    frappe.response["filename"] = filename
    frappe.response["filecontent"] = xlsx.getvalue()
    frappe.response["type"] = "download"
    frappe.response["content_type"] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _lire_xlsx(file_url):
    from frappe.utils.file_manager import get_file_path
    from openpyxl import load_workbook

    filename = get_file_path(file_url)
    workbook = load_workbook(filename=filename, data_only=True)
    sheet = workbook.active
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def _valeur_note(valeur):
    if valeur in (None, ""):
        return None
    try:
        return flt(valeur)
    except Exception:
        return None


@frappe.whitelist()
def export_modele(type_dexamen, cc_columns=None):
    """Télécharge un modèle Excel vierge pour la saisie des notes."""
    cc_columns = _colonnes_cc(cc_columns)
    _repondre_xlsx([_en_tetes(type_dexamen, cc_columns)], "modele_{0}.xlsx".format(type_dexamen.lower()))


def _html_en_pdf(html):
    """Convertit du HTML en PDF (octets).

    Utilise d'abord le générateur standard de Frappe (wkhtmltopdf). S'il
    n'est pas installé, repli sur Playwright/Chromium (déjà présent dans
    l'environnement de développement) pour restituer la même page A4.
    """
    try:
        from frappe.utils.pdf import get_pdf

        return get_pdf(html)
    except Exception:
        pass

    from playwright.sync_api import sync_playwright

    with sync_playwright() as p:
        browser = p.chromium.launch(args=["--no-sandbox"])
        try:
            page = browser.new_page()
            page.set_content(html, wait_until="networkidle")
            pdf = page.pdf(format="A4", landscape=True, print_background=True, margin={"top": "18mm", "bottom": "14mm", "left": "14mm", "right": "14mm"})
            return pdf
        finally:
            browser.close()


def _html_modele_pdf(titre, entetes, lignes):
    """Construit le HTML d'un modèle PDF imprimable — design professionnel."""
    nb_cols = len(entetes)

    # En-têtes du tableau
    entetes_html = ""
    for h in entetes:
        h_esc = frappe.utils.escape_html(h)
        if h in ("Matricule",):
            entetes_html += "<th class='col-mat'>{0}</th>".format(h_esc)
        elif h in ("Nom", "Prénom"):
            entetes_html += "<th class='col-name'>{0}</th>".format(h_esc)
        elif h in ("Crédit",):
            entetes_html += "<th class='col-num'>{0}</th>".format(h_esc)
        else:
            entetes_html += "<th class='col-note'>{0}</th>".format(h_esc)

    # Lignes du tableau
    lignes_html = ""
    for i, ligne in enumerate(lignes):
        row_class = "row-even" if i % 2 == 0 else "row-odd"
        cellules = ""
        for j, c in enumerate(ligne):
            val = frappe.utils.escape_html(str(c) if c is not None else "")
            h = entetes[j] if j < len(entetes) else ""
            if h == "Matricule":
                cellules += "<td class='col-mat'>{0}</td>".format(val)
            elif h in ("Nom", "Prénom"):
                cellules += "<td class='col-name'>{0}</td>".format(val)
            elif h == "Crédit":
                cellules += "<td class='col-num'>{0}</td>".format(val)
            else:
                cellules += "<td class='col-note saisie'></td>"
        lignes_html += "<tr class='{0}'>{1}</tr>".format(row_class, cellules)

    if not lignes_html:
        lignes_html = "<tr><td colspan='{0}' class='empty-row'>Aucun étudiant inscrit</td></tr>".format(nb_cols)

    # Détecter le type depuis le titre pour la couleur d'accent
    titre_lower = titre.lower()
    if "rattrapage" in titre_lower:
        accent = "#e05c2a"
        accent_light = "#fdf0eb"
        badge_label = "RATTRAPAGE"
        badge_bg = "#e05c2a"
    elif "examen" in titre_lower:
        accent = "#2563eb"
        accent_light = "#eff6ff"
        badge_label = "EXAMEN"
        badge_bg = "#2563eb"
    elif "cc" in titre_lower or "contrôle" in titre_lower or "continu" in titre_lower:
        accent = "#7c3aed"
        accent_light = "#f5f3ff"
        badge_label = "CONTRÔLE CONTINU"
        badge_bg = "#7c3aed"
    elif "tp" in titre_lower:
        accent = "#059669"
        accent_light = "#ecfdf5"
        badge_label = "TRAVAUX PRATIQUES"
        badge_bg = "#059669"
    else:
        accent = "#1e40af"
        accent_light = "#eff6ff"
        badge_label = "NOTES"
        badge_bg = "#1e40af"

    return """
<!DOCTYPE html>
<html lang="fr">
<head>
<meta charset="utf-8">
<style>
  @page {{ size: A4 landscape; margin: 14mm 12mm 16mm 12mm; }}
  * {{ box-sizing: border-box; margin: 0; padding: 0; }}
  body {{ font-family: 'Helvetica Neue', Helvetica, Arial, sans-serif; color: #1e293b; background: #fff; font-size: 10px; }}

  /* ── Bandeau supérieur ── */
  .header {{ display: flex; align-items: stretch; margin-bottom: 14px; border-radius: 8px; overflow: hidden; box-shadow: 0 2px 8px rgba(0,0,0,0.10); }}
  .header-accent {{ width: 8px; background: {accent}; flex-shrink: 0; }}
  .header-body {{ flex: 1; padding: 10px 14px; background: #fff; border: 1px solid #e2e8f0; border-left: none; border-radius: 0 8px 8px 0; }}
  .header-top {{ display: flex; align-items: center; justify-content: space-between; margin-bottom: 6px; }}
  .header-title {{ font-size: 15px; font-weight: 700; color: #0f172a; letter-spacing: -0.3px; }}
  .badge {{ display: inline-block; background: {badge_bg}; color: #fff; font-size: 8px; font-weight: 700; letter-spacing: 0.8px; padding: 3px 8px; border-radius: 20px; text-transform: uppercase; }}
  .header-meta {{ display: flex; gap: 18px; }}
  .meta-item {{ font-size: 9px; color: #64748b; }}
  .meta-item strong {{ color: #334155; font-weight: 600; }}
  .header-note {{ font-size: 8.5px; color: #94a3b8; margin-top: 4px; font-style: italic; }}

  /* ── Tableau ── */
  table {{ width: 100%; border-collapse: collapse; font-size: 9.5px; }}
  thead tr {{ background: {accent}; }}
  thead th {{ color: #fff; font-weight: 600; font-size: 8.5px; padding: 7px 8px; text-align: left; letter-spacing: 0.3px; border: 1px solid rgba(255,255,255,0.15); white-space: nowrap; }}
  thead th.col-num, thead th.col-note {{ text-align: center; }}

  tbody tr.row-even {{ background: #fff; }}
  tbody tr.row-odd  {{ background: {accent_light}; }}
  tbody tr:hover    {{ background: #f1f5f9; }}

  td {{ padding: 5px 8px; border: 1px solid #e2e8f0; vertical-align: middle; height: 22px; }}
  td.col-mat  {{ font-family: 'Courier New', monospace; font-size: 8.5px; color: #475569; white-space: nowrap; }}
  td.col-name {{ color: #1e293b; font-weight: 500; }}
  td.col-num  {{ text-align: center; color: #475569; font-weight: 600; }}
  td.col-note {{ text-align: center; min-width: 52px; }}
  td.saisie   {{ background: rgba(255,255,255,0.7); }}

  .empty-row  {{ text-align: center; color: #94a3b8; padding: 20px; font-style: italic; }}

  /* ── Numérotation des lignes ── */
  tbody td:first-child {{ position: relative; }}
  .row-num {{ display: inline-block; width: 14px; height: 14px; line-height: 14px; text-align: center; background: {accent}; color: #fff; border-radius: 50%; font-size: 7px; font-weight: 700; margin-right: 5px; vertical-align: middle; }}

  /* ── Pied de page ── */
  .footer {{ margin-top: 12px; display: flex; justify-content: space-between; align-items: flex-end; }}
  .footer-left {{ font-size: 8px; color: #94a3b8; }}
  .footer-right {{ font-size: 8px; color: #94a3b8; text-align: right; }}
  .signature-block {{ display: flex; gap: 40px; margin-top: 10px; }}
  .signature-item {{ text-align: center; }}
  .signature-line {{ width: 120px; border-bottom: 1px solid #cbd5e1; margin-bottom: 4px; height: 28px; }}
  .signature-label {{ font-size: 8px; color: #64748b; }}

  /* ── Séparateur de section ── */
  .section-bar {{ height: 3px; background: linear-gradient(90deg, {accent} 0%, {accent_light} 100%); border-radius: 2px; margin-bottom: 10px; }}
</style>
</head>
<body>

<div class="header">
  <div class="header-accent"></div>
  <div class="header-body">
    <div class="header-top">
      <div class="header-title">{titre}</div>
      <span class="badge">{badge_label}</span>
    </div>
    <div class="header-meta">
      <div class="meta-item"><strong>Établissement :</strong> UDSHED</div>
      <div class="meta-item"><strong>Barème :</strong> Notes sur 20</div>
      <div class="meta-item"><strong>Étudiants :</strong> {nb_etudiants}</div>
      <div class="meta-item"><strong>Colonnes :</strong> {nb_notes_cols}</div>
    </div>
    <div class="header-note">&#9432; Feuille de saisie officielle — à remplir au stylo et à conserver dans les archives pédagogiques.</div>
  </div>
</div>

<div class="section-bar"></div>

<table>
  <thead><tr>{entetes_html}</tr></thead>
  <tbody>{lignes_html}</tbody>
</table>

<div class="footer">
  <div class="footer-left">
    Document généré par UDSHED &mdash; Confidentiel<br>
    <em>Toute modification doit être paraphée par l'enseignant responsable.</em>
  </div>
  <div class="footer-right">
    <div class="signature-block">
      <div class="signature-item">
        <div class="signature-line"></div>
        <div class="signature-label">Enseignant responsable</div>
      </div>
      <div class="signature-item">
        <div class="signature-line"></div>
        <div class="signature-label">Responsable pédagogique</div>
      </div>
      <div class="signature-item">
        <div class="signature-line"></div>
        <div class="signature-label">Cachet &amp; Date</div>
      </div>
    </div>
  </div>
</div>

</body>
</html>
""".format(
        accent=accent,
        accent_light=accent_light,
        badge_bg=badge_bg,
        badge_label=badge_label,
        titre=frappe.utils.escape_html(titre),
        nb_etudiants=len(lignes),
        nb_notes_cols=sum(1 for h in entetes if h not in ("Matricule", "Nom", "Prénom", "Crédit")),
        entetes_html=entetes_html,
        lignes_html=lignes_html,
    )


def _fmt(v):
    if v is None or v == "":
        return ""
    return str(flt(v))


def _fmt_pct(v):
    if v is None or v == "":
        return ""
    return "{0} %".format(flt(v))


def _code_classe(fos_code, niveau_label):
    """Code de classe compact pour le bandeau du PV (ex : IRT + Licence 3 -> IRT3)."""
    match = re.search(r"(\d+)", niveau_label or "")
    if match and fos_code:
        return "{0}{1}".format(fos_code, match.group(1))
    return niveau_label or ""


def _code_semestre(semestre_label):
    """Code semestre compact pour le bandeau du PV (ex : Semestre 6 -> SEM6)."""
    match = re.search(r"(\d+)", semestre_label or "")
    if match:
        return "SEM{0}".format(match.group(1))
    return semestre_label or ""


def _html_pv_matiere_pdf(
    ecole,
    institut,
    departement,
    logo_html,
    classe,
    annee,
    semestre_label,
    matiere,
    code_matiere,
    inscrits,
    admis,
    echecs,
    taux,
    evaluations,
    lignes,
):
    """Construit le HTML du procès-verbal de la matière (modèle UDM, A4 paysage).

    Toutes les valeurs affichées proviennent du moteur central
    (Session Examen Note) : composantes CC/CCTP/EXAMTP/EXAM, moyenne %,
    grade et points. Aucune donnée n'est recalculée ici.
    """

    def esc(v):
        return frappe.utils.escape_html(str(v) if v is not None else "")

    def note(v):
        return esc(_fmt(v))

    eval_rows = ""
    for ev in evaluations or []:
        eval_rows += (
            "<tr>"
            "<td>{0}</td><td>{1}</td><td>{2}</td><td>{3}</td>"
            "</tr>"
        ).format(esc(ev.get("type")), esc(ev.get("code")), esc(ev.get("matiere")), esc(ev.get("date")))

    notes_rows = ""
    for i, ligne in enumerate(lignes or [], start=1):
        nom_complet = "{0} {1}".format(ligne.get("nom") or "", ligne.get("prenom") or "").strip()
        pct = ligne.get("note_pct")
        notes_rows += (
            "<tr>"
            "<td>{0}</td>"
            "<td>{1}</td>"
            "<td class='nom-cell'>{2}</td>"
            "<td>{3}</td>"
            "<td>{4}</td>"
            "<td>{5}</td>"
            "<td>{6}</td>"
            "<td>{7}</td>"
            "<td>{8}</td>"
            "<td>{9}</td>"
            "</tr>"
        ).format(
            i,
            esc(ligne.get("matricule")),
            esc(nom_complet),
            note(ligne.get("cc")),
            note(ligne.get("cctp")),
            note(ligne.get("examtp")),
            note(ligne.get("examen")),
            esc("{0:.2f}%".format(flt(pct))) if pct is not None else "",
            esc(ligne.get("grade")),
            note(ligne.get("point")),
        )

    if not notes_rows:
        notes_rows = "<tr><td colspan='10' style='padding:12px;color:#666;'>Aucun étudiant inscrit</td></tr>"

    titre_matiere = esc(matiere or "")
    if code_matiere:
        titre_matiere += "({0})".format(esc(code_matiere))

    return """
<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<style>
    * {{ box-sizing: border-box; }}

    body {{
        margin: 0;
        padding: 0;
        background: #ffffff;
        font-family: "Times New Roman", Times, serif;
        color: #000;
    }}

    @page {{ size: A4 landscape; margin: 10mm 12mm; }}

    .page {{ position: relative; }}

    /* ================= HEADER ================= */
    .header {{
        border-bottom: 2px solid #999;
        position: relative;
        margin-bottom: 30px;
        padding-bottom: 8px;
        min-height: 85px;
    }}

    .universite {{
        position: absolute;
        left: 0;
        top: 0;
        color: #071c9d;
        line-height: 1.1;
    }}

    .universite .nom {{
        font-family: Arial, Helvetica, sans-serif;
        font-size: 21px;
        font-weight: bold;
    }}

    .universite .institut {{
        font-family: Arial, Helvetica, sans-serif;
        font-size: 14px;
        margin-top: 4px;
    }}

    .universite .departement {{
        font-family: Arial, Helvetica, sans-serif;
        font-size: 13px;
        margin-top: 3px;
    }}

    .logo {{
        position: absolute;
        right: 0;
        top: -5px;
        width: 110px;
        height: 95px;
        display: flex;
        align-items: center;
        justify-content: center;
    }}
    .logo img {{ max-width: 105px; max-height: 90px; }}

    /* ================= INFORMATIONS ================= */
    .info-header {{
        width: 100%;
        border-collapse: collapse;
        table-layout: fixed;
        margin-bottom: 0;
    }}

    .info-header td {{
        border: 2px solid #222;
        background: #dce2e8;
        text-align: center;
        vertical-align: middle;
        height: 52px;
        font-size: 19px;
        font-weight: bold;
    }}

    .info-header .annee {{ line-height: 1; }}
    .info-header .annee span {{ display: block; font-size: 19px; margin-top: 3px; }}

    /* ================= TITRE ================= */
    .titre {{
        border-left: 2px solid #222;
        border-right: 2px solid #222;
        border-bottom: 2px solid #222;
        text-align: center;
        height: 32px;
        font-size: 14px;
        font-weight: bold;
        padding-top: 6px;
    }}
    .titre span {{ font-weight: normal; }}

    /* ================= STATISTIQUES ================= */
    .stats-zone {{ display: flex; width: 100%; margin-top: 0; }}

    .credits {{ width: 58%; padding-top: 12px; }}
    .evaluations {{ width: 42%; padding-top: 0; }}

    .credits-title {{
        width: 145px;
        height: 20px;
        border: 1px solid #333;
        background: #eef1f3;
        text-align: center;
        color: #d00000;
        font-size: 12px;
        font-weight: bold;
        padding-top: 2px;
    }}

    .credits-table {{ border-collapse: collapse; width: 100%; table-layout: fixed; }}
    .credits-table th,
    .credits-table td {{
        border: 1px solid #555;
        text-align: center;
        height: 20px;
        font-size: 11px;
    }}
    .credits-table th {{ background: #eef1f3; color: #174a8b; font-weight: bold; }}
    .credits-table td {{ color: #174a8b; font-weight: bold; }}

    .evaluations-table {{ width: 100%; border-collapse: collapse; table-layout: fixed; }}
    .evaluations-table th,
    .evaluations-table td {{
        border: 1px solid #444;
        text-align: center;
        height: 19px;
        font-size: 10px;
    }}
    .evaluations-table .title {{
        color: #174a8b;
        background: #eef1f3;
        font-weight: bold;
        height: 18px;
    }}
    .evaluations-table th {{ background: #eef1f3; color: #174a8b; font-weight: bold; }}
    .evaluations-table td {{ color: #174a8b; }}

    /* ================= TABLEAU PRINCIPAL ================= */
    .notes-table {{
        width: 100%;
        border-collapse: collapse;
        table-layout: fixed;
        margin-top: 2px;
    }}

    .notes-table th,
    .notes-table td {{
        border: 1px solid #222;
        text-align: center;
        vertical-align: middle;
    }}

    .notes-table th {{
        height: 26px;
        color: #15519b;
        font-size: 12px;
        font-weight: bold;
        background: #fff;
    }}

    .notes-table td {{ height: 26px; font-size: 12px; }}

    .notes-table .numero {{ width: 42px; }}
    .notes-table .matricule {{ width: 85px; }}
    .notes-table .nom {{ width: 340px; }}
    .notes-table .note {{ width: 52px; }}
    .notes-table .moy {{ width: 58px; }}
    .notes-table .grd {{ width: 48px; }}
    .notes-table .pts {{ width: 50px; }}

    .nom-cell {{ text-align: left !important; padding-left: 8px; white-space: nowrap; }}

    /* ================= SIGNATURES ================= */
    .signatures {{
        width: 100%;
        display: flex;
        justify-content: space-between;
        margin-top: 26px;
    }}

    .signature {{
        width: 23%;
        height: 30px;
        border: 2px solid #222;
        text-align: center;
        font-size: 13px;
        font-weight: bold;
        padding-top: 6px;
    }}

    /* ================= FILIGRANE ================= */
    .watermark {{
        position: fixed;
        left: 50%;
        top: 55%;
        transform: translate(-50%, -50%) rotate(-28deg);
        font-family: Arial, Helvetica, sans-serif;
        font-size: 52px;
        font-weight: bold;
        color: rgba(130, 130, 130, 0.13);
        white-space: nowrap;
        pointer-events: none;
        z-index: 10;
    }}
</style>
</head>
<body>

<div class="page">

    <div class="watermark">{0}</div>

    <div class="header">
        <div class="universite">
            <div class="nom">{0}</div>
            <div class="institut">{1}</div>
            <div class="departement">{2}</div>
        </div>
        <div class="logo">{3}</div>
    </div>

    <table class="info-header">
        <tr>
            <td>{4}</td>
            <td class="annee">
                Année Académique
                <span>{5}</span>
            </td>
            <td>{6}</td>
        </tr>
    </table>

    <div class="titre">
        <b>PROCES VERBAL DE LA MATIERE:</b>
        <span>{7}</span>
    </div>

    <div class="stats-zone">
        <div class="credits">
            <div class="credits-title">Credits:</div>
            <table class="credits-table">
                <tr>
                    <th>Inscrits</th>
                    <th>Admis</th>
                    <th>Echecs</th>
                    <th>Taux de réussite</th>
                </tr>
                <tr>
                    <td>{8}</td>
                    <td>{9}</td>
                    <td>{10}</td>
                    <td>{11}</td>
                </tr>
            </table>
        </div>

        <div class="evaluations">
            <table class="evaluations-table">
                <tr>
                    <td colspan="4" class="title">Evaluations effectuées:</td>
                </tr>
                <tr>
                    <th>Type</th>
                    <th>Code</th>
                    <th>Matière</th>
                    <th>Date</th>
                </tr>
                {12}
            </table>
        </div>
    </div>

    <table class="notes-table">
        <thead>
            <tr>
                <th class="numero">N°</th>
                <th class="matricule">Matricule</th>
                <th class="nom">Nom et Prénoms</th>
                <th class="note">CC</th>
                <th class="note">CCTP</th>
                <th class="note">EXAMTP</th>
                <th class="note">EXAM</th>
                <th class="moy">MOY</th>
                <th class="grd">GRD</th>
                <th class="pts">PTS</th>
            </tr>
        </thead>
        <tbody>
            {13}
        </tbody>
    </table>

    <div class="signatures">
        <div class="signature">VICE DOYEN</div>
        <div class="signature">COORDONATEUR</div>
        <div class="signature">RESP ACADEMIQUE</div>
        <div class="signature">CHEF UNITE</div>
    </div>

</div>

</body>
</html>
    """.format(
        esc(ecole),
        esc(institut),
        esc(departement),
        logo_html,
        esc(classe),
        esc(annee),
        esc(semestre_label),
        titre_matiere,
        inscrits,
        admis,
        echecs,
        taux,
        eval_rows,
        notes_rows,
    )


@frappe.whitelist()
def generer_pdf(academic_year, filiere, niveau, semestre, teaching_unit):
    """Génère le procès-verbal de la matière au format PDF à partir des données enregistrées.

    Le PV présente, pour chaque étudiant inscrit : les composantes saisies
    (CC, CCTP, EXAMTP, EXAM), la moyenne (%), le grade et les points, ainsi
    que les statistiques de l'UE (inscrits, admis, échecs, taux de réussite)
    et les évaluations de la session. Toutes les valeurs proviennent du
    moteur central (Session Examen Note) — aucune donnée fictive, aucun
    recalcul local. Si un rattrapage est saisi, la note retenue (= MAX
    examen/rattrapage) et le résultat correspondant sont utilisés.
    """
    _verifier_acces_enseignant(teaching_unit)
    args = {
        "academic_year": academic_year,
        "semestre": _semestre_effectif({"academic_year": academic_year, "semestre": semestre, "teaching_unit": teaching_unit}),
        "filiere": filiere,
        "niveau": niveau,
        "teaching_unit": teaching_unit,
    }
    semestre_effectif = args["semestre"]

    ue_info = _get_ue_info(teaching_unit, filiere, niveau)
    students = _get_etudiants(academic_year, filiere, niveau, teaching_unit)

    sessions = {
        "cc": _get_or_create_session(args, TYPE_CC),
        "normale": _get_or_create_session(args, TYPE_NORMALE),
        "rattrapage": _get_or_create_session(args, TYPE_RATTRAPAGE),
    }
    notes = _charger_notes(students, teaching_unit, sessions)

    # Lignes unifiées (CC, CCTP, EXAMTP, EXAM + résultat) — même source que la page.
    lignes = _lignes_unifiees(students, notes)

    # Le rattrapage saisi prime : note retenue et résultat de la session de rattrapage.
    notes_rt = notes.get("Rattrapage", {})
    for ligne in lignes:
        rt = notes_rt.get(ligne.get("student")) or {}
        if rt.get("note_pct") is not None:
            ligne["examen"] = rt.get("note_examen_active")
            ligne["note_finale"] = rt.get("note_finale")
            ligne["note_pct"] = rt.get("note_pct")
            ligne["grade"] = rt.get("grade")
            ligne["point"] = rt.get("point")

    # Statistiques de l'UE : admis/échecs selon le seuil de validation du cycle.
    cycle = _cycle_pour_niveau(niveau)
    seuil = get_seuil_validation(cycle)
    evalues = [l for l in lignes if l.get("note_pct") is not None]
    admis = sum(1 for l in evalues if flt(l["note_pct"]) >= seuil)
    echecs = len(evalues) - admis
    taux = "{0:.2f}%".format(admis * 100.0 / len(evalues)) if evalues else ""

    # Évaluations de la session (type + matière + date de début).
    code_matiere = ue_info.get("code") or ""
    evaluations = []
    for type_label, cle in (("CC", "cc"), ("EXAM", "normale"), ("RAT", "rattrapage")):
        date_debut = frappe.db.get_value("Session Examen", sessions[cle], "date_debut")
        evaluations.append(
            {
                "type": type_label,
                "code": "",
                "matiere": code_matiere,
                "date": frappe.utils.formatdate(date_debut, "dd-MM-yyyy") if date_debut else "",
            }
        )

    setting = frappe.get_single("Udshed Setting")
    ecole = setting.school_name or ""

    logo_html = ""
    if setting.school_logo:
        logo_url = setting.school_logo
        if logo_url.startswith("/"):
            logo_url = frappe.utils.get_url(logo_url)
        logo_html = '<img src="{0}" alt="logo">'.format(frappe.utils.escape_html(logo_url))

    fos_doc = frappe.get_doc("Field of study", filiere)
    faculte = ""
    if fos_doc.faculte:
        faculte = frappe.db.get_value("Faculty", fos_doc.faculte, "faculty_name") or ""

    html = _html_pv_matiere_pdf(
        ecole=ecole,
        institut=faculte,
        departement=fos_doc.name_of_field or "",
        logo_html=logo_html,
        classe=_code_classe(fos_doc.field_of_study_code or "", niveau),
        annee=academic_year,
        semestre_label=_code_semestre(semestre_effectif),
        matiere=ue_info.get("intitule") or teaching_unit,
        code_matiere=code_matiere,
        inscrits=len(lignes),
        admis=admis,
        echecs=echecs,
        taux=taux,
        evaluations=evaluations,
        lignes=lignes,
    )

    frappe.response["filename"] = "PV_{0}_{1}.pdf".format(
        (code_matiere or teaching_unit).replace("/", "-"),
        (niveau or "").replace(" ", "-"),
    )
    frappe.response["filecontent"] = _html_en_pdf(html)
    frappe.response["type"] = "download"
    frappe.response["content_type"] = "application/pdf"


@frappe.whitelist()
def export_modele_pdf(
    type_dexamen,
    cc_columns=None,
    academic_year=None,
    filiere=None,
    niveau=None,
    semestre=None,
    teaching_unit=None,
):
    """Télécharge un modèle PDF imprimable, étudiants pré-remplis si un contexte est fourni."""
    cc_columns = _colonnes_cc(cc_columns)
    entetes = _en_tetes(type_dexamen, cc_columns)
    if teaching_unit:
        entetes = _entetes_avec_tp(entetes, type_dexamen, teaching_unit)

    students = []
    credit = ""
    if teaching_unit and academic_year and filiere and niveau:
        students = _get_etudiants(academic_year, filiere, niveau, teaching_unit)
        credit = _credits_ue(teaching_unit, filiere, niveau)

    lignes = []
    for s in students:
        cellule = []
        for h in entetes:
            if h == "Matricule":
                cellule.append(s.get("matricule") or s.get("student") or "")
            elif h == "Nom":
                cellule.append(s.get("nom") or "")
            elif h == "Prénom":
                cellule.append(s.get("prenom") or "")
            elif h == "Crédit":
                cellule.append(credit)
            else:
                cellule.append("")
        lignes.append(cellule)

    titre = _("Modèle de saisie des notes — {0}").format(type_dexamen)
    html = _html_modele_pdf(titre, entetes, lignes)

    frappe.response["filename"] = "modele_{0}.pdf".format(type_dexamen.lower())
    frappe.response["filecontent"] = _html_en_pdf(html)
    frappe.response["type"] = "download"
    frappe.response["content_type"] = "application/pdf"


@frappe.whitelist()
def export_notes(academic_year, filiere, niveau, semestre, teaching_unit, type_dexamen, cc_columns=None):
    """Exporte les notes déjà saisies au format Excel."""
    cc_columns = _colonnes_cc(cc_columns)
    args = {"academic_year": academic_year, "filiere": filiere, "niveau": niveau, "semestre": _semestre_effectif({"academic_year": academic_year, "semestre": semestre, "teaching_unit": teaching_unit}), "teaching_unit": teaching_unit}
    students = _get_etudiants(academic_year, filiere, niveau, teaching_unit)

    sessions = {
        "cc": _get_or_create_session(args, TYPE_CC),
        "normale": _get_or_create_session(args, TYPE_NORMALE),
        "rattrapage": _get_or_create_session(args, TYPE_RATTRAPAGE),
    }
    notes = _charger_notes(students, teaching_unit, sessions)
    credit = _credits_ue(teaching_unit, filiere, niveau)

    entetes = _en_tetes(type_dexamen, cc_columns)
    entetes = _entetes_avec_tp(entetes, type_dexamen, teaching_unit)
    lignes = [entetes]
    for s in students:
        nom = s["student"]
        if type_dexamen == "CC":
            cc = notes["CC"].get(nom, {})
            valeurs = {}
            for item in cc.get("notes_cc", []):
                valeurs[item.get("cc_label")] = item.get("note_cc")
            lignes.append(
                [s["matricule"], s["nom"], s["prenom"], credit]
                + [valeurs.get(c.get("label")) for c in cc_columns]
                + [cc.get("note_cc_moyenne")]
            )
        elif type_dexamen == "Examen":
            ex = notes["Examen"].get(nom, {})
            ligne = [s["matricule"], s["nom"], s["prenom"], credit]
            if _("Note de TP") in entetes:
                ligne.append(ex.get("note_tp"))
            ligne.append(ex.get("note_examen"))
            lignes.append(ligne)
        elif type_dexamen == "TP":
            tp = notes["TP"].get(nom, {})
            lignes.append([s["matricule"], s["nom"], s["prenom"], credit, tp.get("note_tp")])
        elif type_dexamen == "Rattrapage":
            rt = notes["Rattrapage"].get(nom, {})
            lignes.append(
                [s["matricule"], s["nom"], s["prenom"], rt.get("note_examen"), rt.get("note_examen_rattrapage"), rt.get("note_examen_active")]
            )

    _repondre_xlsx(lignes, "notes_{0}_{1}.xlsx".format(teaching_unit.replace("/", "-"), type_dexamen.lower()))


@frappe.whitelist()
def importer_notes(file_url, type_dexamen, academic_year, filiere, niveau, semestre, teaching_unit, cc_columns=None):
    """Importe des notes depuis un fichier Excel (modèle exporté)."""
    cc_columns = _colonnes_cc(cc_columns)
    args = {"academic_year": academic_year, "filiere": filiere, "niveau": niveau, "semestre": _semestre_effectif({"academic_year": academic_year, "semestre": semestre, "teaching_unit": teaching_unit}), "teaching_unit": teaching_unit}
    students = _get_etudiants(academic_year, filiere, niveau, teaching_unit)
    par_matricule = {s["matricule"]: s["student"] for s in students}

    lignes = _lire_xlsx(file_url)
    if not lignes or not lignes[0]:
        frappe.throw(_("Le fichier Excel est vide."))

    entetes = [str(h).strip() if h is not None else "" for h in lignes[0]]
    idx_matricule = entetes.index("Matricule") if "Matricule" in entetes else 0

    erreurs = []
    a_sauver = []

    for ligne in lignes[1:]:
        if not ligne or not any(l not in (None, "") for l in ligne):
            continue
        matricule = str(ligne[idx_matricule]).strip() if idx_matricule < len(ligne) and ligne[idx_matricule] is not None else ""
        student = par_matricule.get(matricule)
        if not student:
            erreurs.append(_("Matricule inconnu : {0}").format(matricule))
            continue

        if type_dexamen == "CC":
            items = []
            for i, c in enumerate(cc_columns):
                col = entetes.index(c.get("label")) if c.get("label") in entetes else None
                if col is None:
                    continue
                valeur = _valeur_note(ligne[col]) if col < len(ligne) else None
                if valeur is None:
                    continue
                items.append({"cc_label": c.get("label"), "cc_weight": c.get("weight"), "note_cc": valeur})
            if items:
                a_sauver.append({"student": student, "notes_cc": items})
        else:
            col_note = {
                "Examen": "Note d'examen",
                "TP": "Note de TP",
                "Rattrapage": "Note de rattrapage",
            }[type_dexamen]
            idx = entetes.index(col_note) if col_note in entetes else None
            valeur = _valeur_note(ligne[idx]) if idx is not None and idx < len(ligne) else None
            if valeur is None:
                continue
            champ = {"Examen": "note_examen", "TP": "note_tp", "Rattrapage": "note_examen_rattrapage"}[type_dexamen]
            a_sauver.append({"student": student, champ: valeur})

    if erreurs:
        frappe.throw(_("Import impossible :\n- {0}").format("\n- ".join(erreurs[:50])))

    if type_dexamen == "CC":
        n = _sauvegarder_cc(args, a_sauver)
    elif type_dexamen == "Examen":
        n = _sauvegarder_examen(args, a_sauver)
    elif type_dexamen == "TP":
        n = _sauvegarder_tp(args, a_sauver)
    else:
        n = _sauvegarder_rattrapage(args, a_sauver)

    return {"saved": n}


@frappe.whitelist()
def export_evaluations(academic_year, filiere, niveau, semestre, teaching_unit):
    """Exporte la saisie unifiée (CC, CCTP, EXAMTP, EXAM) au format Excel."""
    args = {"academic_year": academic_year, "filiere": filiere, "niveau": niveau, "semestre": _semestre_effectif({"academic_year": academic_year, "semestre": semestre, "teaching_unit": teaching_unit}), "teaching_unit": teaching_unit}
    students = _get_etudiants(academic_year, filiere, niveau, teaching_unit)
    sessions = {
        "cc": _get_or_create_session(args, TYPE_CC),
        "normale": _get_or_create_session(args, TYPE_NORMALE),
        "rattrapage": _get_or_create_session(args, TYPE_RATTRAPAGE),
    }
    notes = _charger_notes(students, teaching_unit, sessions)
    credit = _credits_ue(teaching_unit, filiere, niveau)

    lignes = [["Matricule", "Nom", "Prénoms", "Crédits", "CC", "CCTP", "EXAMTP", "EXAM", "MOY", "GRD", "PTS"]]
    for ligne in _lignes_unifiees(students, notes):
        lignes.append(
            [
                ligne["matricule"],
                ligne["nom"],
                ligne["prenom"],
                credit,
                ligne.get("cc"),
                ligne.get("cctp"),
                ligne.get("examtp"),
                ligne.get("examen"),
                ligne.get("note_finale"),
                ligne.get("grade"),
                ligne.get("point"),
            ]
        )

    _repondre_xlsx(lignes, "notes_{0}_evaluations.xlsx".format(teaching_unit.replace("/", "-")))


@frappe.whitelist()
def importer_evaluations(file_url, academic_year, filiere, niveau, semestre, teaching_unit):
    """Importe la saisie unifiée depuis un Excel (Matricule + CC + CCTP + EXAMTP + EXAM)."""
    args = {"academic_year": academic_year, "filiere": filiere, "niveau": niveau, "semestre": _semestre_effectif({"academic_year": academic_year, "semestre": semestre, "teaching_unit": teaching_unit}), "teaching_unit": teaching_unit}
    students = _get_etudiants(academic_year, filiere, niveau, teaching_unit)
    par_matricule = {s["matricule"]: s["student"] for s in students}

    lignes = _lire_xlsx(file_url)
    if not lignes or not lignes[0]:
        frappe.throw(_("Le fichier Excel est vide."))

    entetes = [str(h).strip() if h is not None else "" for h in lignes[0]]
    idx_matricule = entetes.index("Matricule") if "Matricule" in entetes else 0

    def _idx_col(nom):
        return entetes.index(nom) if nom in entetes else None

    erreurs = []
    a_sauver = []
    for ligne in lignes[1:]:
        if not ligne or not any(l not in (None, "") for l in ligne):
            continue
        matricule = str(ligne[idx_matricule]).strip() if idx_matricule < len(ligne) and ligne[idx_matricule] is not None else ""
        student = par_matricule.get(matricule)
        if not student:
            erreurs.append(_("Matricule inconnu : {0}").format(matricule))
            continue
        row = {"student": student}
        for champ, colonne in (("cc", "CC"), ("note_cctp", "CCTP"), ("note_examtp", "EXAMTP"), ("note_examen", "EXAM")):
            i = _idx_col(colonne)
            valeur = _valeur_note(ligne[i]) if i is not None and i < len(ligne) else None
            if valeur is not None:
                row[champ] = valeur
        if any(row.get(c) is not None for c in ("cc", "note_cctp", "note_examtp", "note_examen")):
            a_sauver.append(row)

    if erreurs:
        frappe.throw(_("Import impossible :\n- {0}").format("\n- ".join(erreurs[:50])))

    n = 0
    if a_sauver:
        n = enregistrer_evaluations(academic_year, filiere, niveau, semestre, teaching_unit, a_sauver)["saved"]
    return {"saved": n}
