# Copyright (c) 2026, Cédric Nguendap Bedjama and contributors
# For license information, please see license.txt

"""Tests des inscriptions hors candidature (page « Liste des inscriptions »).

Couvre :
  - import Excel et CSV (simulation puis validation) ;
  - matricules générés (série Z) et matricules existants conservés ;
  - doublons (fichier, matricule, e-mail, identité), dossiers réutilisés,
    inscriptions déjà présentes ou d'une autre année ;
  - bilan (importées / ignorées / rejetées + motif) ;
  - export Excel et CSV filtré, sans écriture ;
  - parcours « arrivée en cours d'année » : matricule généré au
    téléchargement de la fiche, une seule fois ;
  - permissions.
"""

import csv
import io
import re
from unittest.mock import patch

import frappe
from frappe.tests import IntegrationTestCase
from openpyxl import load_workbook

from udshed.api import inscription_import_export as module
from udshed.api.inscription_import_export import (
    _lire_fichier,
    _traiter_import,
    exporter_inscriptions,
    generer_matricule_et_fiche,
    importer_inscriptions,
    lister_inscriptions,
    telecharger_modele_import,
)
from udshed.udshed._fixture_factory import (
    _COMPTEUR,
    make_academic_year,
    make_course,
    make_faculty,
    make_field_of_study,
    make_level,
    make_student,
    make_teaching_unit,
    make_user_with_role,
    suppress_commits,
)

ENTETES = [
    "Matricule", "Nom", "Prénom", "Sexe", "Date de naissance", "Lieu de naissance",
    "Email", "Téléphone", "Filière", "Niveau", "Année académique",
]


class TestImportExportInscriptions(IntegrationTestCase):
    def setUp(self):
        frappe.set_user("Administrator")
        self.annee = make_academic_year("2026", "2027")
        self.annee_precedente = make_academic_year("2025", "2026")
        frappe.db.set_single_value("Udshed Setting", "current_year", self.annee.name)

        self.fos = make_field_of_study(make_faculty())
        self.niveau = make_level(self.fos, level="Licence 1")
        self.fos_b = make_field_of_study(make_faculty())
        self.niveau_b = make_level(self.fos_b, level="Licence 2", order=2)
        self.tu = make_teaching_unit(make_course(), self.annee, self.fos, self.niveau)

        self.fichiers = []
        patcher = patch.object(module, "_envoyer_email_matricule_async")
        self.email = patcher.start()
        self.addCleanup(patcher.stop)

    def tearDown(self):
        frappe.set_user("Administrator")
        for nom in self.fichiers:
            if frappe.db.exists("File", nom):
                frappe.delete_doc("File", nom, ignore_permissions=True, force=True)

    # ------------------------------------------------------------------ #
    #  Helpers
    # ------------------------------------------------------------------ #
    def _ligne(self, matricule="", nom=None, prenom="Jean", sexe="M", naissance="15/03/2004",
               email=None, filiere=None, niveau="Licence 1", annee=""):
        _COMPTEUR["n"] += 1
        n = _COMPTEUR["n"]
        return [
            matricule, nom or f"IMPORT{n}", prenom, sexe, naissance, "Dschang",
            email if email is not None else f"import{n}@udshed.test", "655000000",
            filiere or self.fos.name, niveau, annee,
        ]

    def _fichier(self, lignes, format="csv"):
        if format == "csv":
            flux = io.StringIO()
            csv.writer(flux, delimiter=";").writerows([ENTETES, *lignes])
            contenu = ("﻿" + flux.getvalue()).encode("utf-8")
        else:
            from frappe.utils.xlsxutils import make_xlsx

            contenu = make_xlsx([ENTETES, *lignes], "Étudiants").getvalue()
        doc = frappe.get_doc({
            "doctype": "File",
            "file_name": f"import_test_{frappe.generate_hash(length=8)}.{format}",
            "content": contenu,
            "is_private": 1,
        }).insert(ignore_permissions=True)
        self.fichiers.append(doc.name)
        return doc.file_url

    def _importer(self, lignes, format="csv", valider=1, annee=None):
        file_url = self._fichier(lignes, format)
        with suppress_commits():
            return importer_inscriptions(file_url, annee_academique=annee or self.annee.name, valider=valider)

    def _bilan(self, lignes):
        return _traiter_import([ENTETES, *lignes], self.annee.name, executer=True)

    def _ia_manuelle(self, **valeurs):
        _COMPTEUR["n"] += 1
        n = _COMPTEUR["n"]
        doc = frappe.get_doc({
            "doctype": "Inscription Academique",
            "nom_prenom": f"MANUEL{n} Paul",
            "email": f"manuel{n}@udshed.test",
            "date_naissance": "2003-05-02",
            "lieu_naissance": "Douala",
            "telephone": "+237677000000",
            "sexe": "Masculin",
            "annee_academique": self.annee.name,
            "filiere": self.fos.name,
            "classe": "Licence 1",
            **valeurs,
        })
        doc.insert()
        return doc

    def _prefixe(self):
        return self.annee.name.split("-")[0][-2:] + "Z"

    # ------------------------------------------------------------------ #
    #  Import
    # ------------------------------------------------------------------ #
    def test_import_csv_simulation_puis_validation(self):
        lignes = [self._ligne(), self._ligne(sexe="F", prenom="Marie")]
        avant = frappe.db.count("Inscription Academique")

        simulation = self._importer(lignes, valider=0)
        self.assertFalse(simulation["valide"])
        self.assertEqual(simulation["importees"], 2)
        self.assertEqual({l["statut"] for l in simulation["lignes"]}, {"À importer"})
        self.assertEqual(frappe.db.count("Inscription Academique"), avant, "la simulation n'écrit rien")

        bilan = self._importer(lignes)
        self.assertTrue(bilan["valide"])
        self.assertEqual((bilan["importees"], bilan["ignorees"], bilan["rejetees"]), (2, 0, 0))

        matricules = [l["matricule"] for l in bilan["lignes"]]
        for m in matricules:
            self.assertRegex(m, rf"^{self._prefixe()}\d{{3}}$")
        numeros = sorted(int(m[3:]) for m in matricules)
        self.assertEqual(numeros[1], numeros[0] + 1, "matricules consécutifs et uniques")

        for m in matricules:
            student = frappe.get_doc("Student", {"matricule": m})
            self.assertTrue(frappe.db.exists("User", student.email), "compte utilisateur créé")
            self.assertEqual(student.utilisateur, student.email)
            self.assertTrue(frappe.db.exists(
                "Academic Reregistration", {"student": student.name, "academic_year": self.annee.name}
            ), "inscrit aux UE")
            ia = frappe.get_doc("Inscription Academique", {"matricule": m})
            self.assertFalse(ia.dossier_origine)
        self.assertEqual(self.email.call_count, 2, "un e-mail par étudiant importé")
        self.assertEqual(frappe.get_doc("Student", {"matricule": matricules[1]}).sexe, "Femme")

    def test_import_xlsx_matricule_existant_conserve(self):
        matricule = f"OLD{frappe.generate_hash(length=6).upper()}"
        bilan = self._importer([self._ligne(matricule=matricule.lower())], format="xlsx")
        self.assertEqual(bilan["importees"], 1, bilan)
        self.assertEqual(bilan["lignes"][0]["matricule"], matricule)
        self.assertTrue(frappe.db.exists("Student", {"matricule": matricule}))
        self.assertTrue(frappe.db.exists("Inscription Academique", {"matricule": matricule}))

    def test_lecture_xlsx_et_csv_identiques(self):
        lignes = [self._ligne()]
        self.assertEqual(
            [str(v or "") for v in _lire_fichier(self._fichier(lignes, "xlsx"))[1]],
            [str(v or "") for v in _lire_fichier(self._fichier(lignes, "csv"))[1]],
        )

    def test_doublons_dans_le_fichier(self):
        premiere = self._ligne()
        meme_email = self._ligne(email=premiere[6])
        meme_identite = self._ligne(nom=premiere[1], prenom=premiere[2])
        bilan = self._bilan([premiere, meme_email, meme_identite])
        self.assertEqual([l["statut"] for l in bilan["lignes"]], ["Importée", "Rejetée", "Rejetée"])
        self.assertIn("Doublon dans le fichier", bilan["lignes"][1]["motif"])
        self.assertIn("ligne 2", bilan["lignes"][1]["motif"])
        self.assertIn("identité", bilan["lignes"][2]["motif"])

    def test_etudiant_existant_reutilise_sans_changer_son_matricule(self):
        student = make_student(self.fos, self.niveau)
        nb_students = frappe.db.count("Student")
        # Même identité (nom, prénom, naissance), e-mail du fichier différent
        ligne = self._ligne(nom=student.nom, prenom=student.prenom, naissance="01/01/2000",
                            email=f"autre{_COMPTEUR['n']}@udshed.test")
        bilan = self._bilan([ligne])
        self.assertEqual(bilan["importees"], 1, bilan)
        self.assertEqual(bilan["lignes"][0]["matricule"], student.matricule)
        self.assertIn("Étudiant existant", bilan["lignes"][0]["motif"])
        self.assertEqual(frappe.db.count("Student"), nb_students, "aucun Student dupliqué")
        self.assertEqual(frappe.db.get_value("Student", student.name, "matricule"), student.matricule)
        self.assertTrue(frappe.db.exists(
            "Inscription Academique", {"matricule": student.matricule, "annee_academique": self.annee.name}
        ))

    def test_reimport_inscription_existante_ignoree(self):
        lignes = [self._ligne(), self._ligne()]
        premier = self._bilan(lignes)
        avant = frappe.db.count("Inscription Academique")
        second = self._bilan(lignes)
        self.assertEqual(second["ignorees"], 2, second)
        self.assertEqual(frappe.db.count("Inscription Academique"), avant)
        self.assertEqual(
            [l["matricule"] for l in second["lignes"]], [l["matricule"] for l in premier["lignes"]],
            "les matricules ne changent pas à la réimportation",
        )
        # Même étudiant avec son matricule explicite : ignoré aussi
        avec_matricule = list(lignes[0])
        avec_matricule[0] = premier["lignes"][0]["matricule"]
        self.assertEqual(self._bilan([avec_matricule])["lignes"][0]["statut"], "Ignorée")

    def test_inscription_autre_annee_renvoie_vers_reinscription(self):
        ligne = self._ligne(annee=self.annee_precedente.name)
        self.assertEqual(self._bilan([ligne])["importees"], 1)
        nouvelle = list(ligne)
        nouvelle[10] = self.annee.name
        bilan = self._bilan([nouvelle])
        self.assertEqual(bilan["lignes"][0]["statut"], "Rejetée")
        self.assertIn("réinscription", bilan["lignes"][0]["motif"])

    def test_conflit_matricule_et_email(self):
        a = make_student(self.fos, self.niveau)
        b = make_student(self.fos, self.niveau)
        bilan = self._bilan([self._ligne(matricule=a.matricule, email=b.email)])
        self.assertEqual(bilan["lignes"][0]["statut"], "Rejetée")
        self.assertIn("Conflit", bilan["lignes"][0]["motif"])

    def test_lignes_invalides_rejetees_avec_motif(self):
        bilan = self._bilan([
            self._ligne(sexe="X"),
            self._ligne(filiere="FILIERE-INCONNUE"),
            self._ligne(niveau="Doctorat 9"),
            self._ligne(email="pas-un-email"),
            self._ligne(naissance="31/02/2004"),
            self._ligne(matricule="!!"),
            self._ligne(annee="1990-1991"),
        ])
        self.assertEqual(bilan["rejetees"], 7)
        motifs = [l["motif"] for l in bilan["lignes"]]
        for attendu, motif in zip(
            ["Sexe invalide", "Filière inconnue", "Niveau", "E-mail", "Date de naissance", "Matricule invalide", "Année académique inconnue"],
            motifs,
        ):
            self.assertIn(attendu, motif)

    def test_colonnes_obligatoires_absentes(self):
        with self.assertRaises(frappe.ValidationError):
            _traiter_import([["Nom", "Prénom"], ["A", "B"]], self.annee.name, executer=False)

    def test_modele_import(self):
        telecharger_modele_import("xlsx")
        wb = load_workbook(io.BytesIO(frappe.response["filecontent"]))
        self.assertEqual(wb.sheetnames, ["Étudiants", "Instructions", "Filières"])
        self.assertEqual([c.value for c in wb["Étudiants"][1]][:len(ENTETES)], ENTETES)
        telecharger_modele_import("csv")
        self.assertTrue(frappe.response["filecontent"].decode("utf-8-sig").startswith("Matricule;Nom;"))

    # ------------------------------------------------------------------ #
    #  Arrivée en cours d'année
    # ------------------------------------------------------------------ #
    def test_matricule_genere_au_telechargement_une_seule_fois(self):
        ia = self._ia_manuelle()
        self.assertFalse(ia.matricule)
        self.assertIn(
            ia.name,
            [r["inscription"] for r in lister_inscriptions(statut=module.STATUT_EN_ATTENTE)["rows"]],
        )

        with suppress_commits():
            premier = generer_matricule_et_fiche(ia.name)
            second = generer_matricule_et_fiche(ia.name)
        self.assertRegex(premier["matricule"], rf"^{self._prefixe()}\d{{3}}$")
        self.assertEqual(second["matricule"], premier["matricule"])
        self.assertIn("telecharger_fiche_inscription?nom=", premier["pdf_url"])
        self.assertEqual(self.email.call_count, 1, "pas de second e-mail")
        student = frappe.get_doc("Student", {"matricule": premier["matricule"]})
        self.assertEqual((student.nom, student.prenom), (ia.nom_prenom.split()[0], "Paul"))
        self.assertTrue(frappe.db.exists("Academic Reregistration", {"student": student.name}))

    def test_saisie_manuelle_niveau_invalide(self):
        with self.assertRaises(frappe.ValidationError):
            self._ia_manuelle(classe="Niveau inexistant")

    # ------------------------------------------------------------------ #
    #  Export
    # ------------------------------------------------------------------ #
    def test_export_csv_et_xlsx_respectent_les_filtres(self):
        self._bilan([self._ligne(), self._ligne(filiere=self.fos_b.name, niveau="Licence 2")])
        avant = frappe.db.sql("SELECT name, matricule FROM `tabInscription Academique` ORDER BY name")

        exporter_inscriptions(format="csv", filiere=self.fos.name)
        lignes = list(csv.reader(io.StringIO(frappe.response["filecontent"].decode("utf-8-sig")), delimiter=";"))
        self.assertEqual(lignes[0][0], "Matricule")
        self.assertIn("Statut de l'inscription", lignes[0])
        self.assertTrue(len(lignes) > 1)
        idx = lignes[0].index("Code filière")
        self.assertEqual({l[idx] for l in lignes[1:]}, {self.fos.name})

        exporter_inscriptions(format="xlsx", filiere=self.fos_b.name, niveau="Licence 2")
        ws = load_workbook(io.BytesIO(frappe.response["filecontent"])).active
        rows = list(ws.iter_rows(values_only=True))
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][rows[0].index("Code filière")], self.fos_b.name)

        exporter_inscriptions(format="csv", filiere=self.fos.name, recherche="aucun-resultat-xyz")
        self.assertEqual(len(frappe.response["filecontent"].decode("utf-8-sig").strip().splitlines()), 1)

        self.assertEqual(
            frappe.db.sql("SELECT name, matricule FROM `tabInscription Academique` ORDER BY name"), avant,
            "l'export ne crée ni ne modifie rien",
        )

    # ------------------------------------------------------------------ #
    #  Permissions
    # ------------------------------------------------------------------ #
    def test_permissions(self):
        sans_droit = make_user_with_role("Teacher")
        frappe.set_user(sans_droit.name)
        with self.assertRaises(frappe.PermissionError):
            exporter_inscriptions(format="csv")
        with self.assertRaises(frappe.PermissionError):
            importer_inscriptions("/private/files/inexistant.csv")
        with self.assertRaises(frappe.PermissionError):
            telecharger_modele_import()

        # Comptable : lecture/export oui, import non
        frappe.set_user("Administrator")
        comptable = make_user_with_role("Comptable")
        frappe.set_user(comptable.name)
        exporter_inscriptions(format="csv")
        with self.assertRaises(frappe.PermissionError):
            importer_inscriptions("/private/files/inexistant.csv")

        # Registration Manager : import autorisé
        frappe.set_user("Administrator")
        responsable = make_user_with_role("Registration Manager")
        file_url = self._fichier([self._ligne()])
        frappe.db.set_value("File", {"file_url": file_url}, "owner", responsable.name)
        frappe.set_user(responsable.name)
        with suppress_commits():
            bilan = importer_inscriptions(file_url, annee_academique=self.annee.name, valider=1)
        self.assertEqual(bilan["importees"], 1, bilan)
        self.assertTrue(re.match(rf"^{self._prefixe()}\d{{3}}$", bilan["lignes"][0]["matricule"]))
